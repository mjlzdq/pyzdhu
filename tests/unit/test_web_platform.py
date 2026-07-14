"""
Web 平台后端单元测试

覆盖 web_platform/backend/core 下不依赖网络的纯逻辑：
- parser：CSV / XLSX 解析、JSON 自动识别、空行过滤
- renderer：动态占位符与变量池渲染
- runner：请求构建、变量提取、断言逻辑
"""
import io
import json

import openpyxl
import pytest

from web_platform.backend.core.parser import (
    parse_csv, parse_xlsx, parse_file, detect_format, _auto_parse_json,
)
from web_platform.backend.core.renderer import (
    _render_dynamic_placeholders, _render_all_placeholders,
)
from web_platform.backend.core.runner import TestRunner, _truncate


pytestmark = pytest.mark.unit


# ==================== parser ====================

class TestParserCSV:
    def test_parse_basic_csv(self):
        content = "case_id,method,path\nC1,GET,/posts\nC2,POST,/posts\n"
        rows = parse_csv(content)
        assert len(rows) == 2
        assert rows[0]["case_id"] == "C1"
        assert rows[1]["method"] == "POST"

    def test_strip_whitespace(self):
        content = "case_id, method \nC1,  GET  \n"
        rows = parse_csv(content)
        assert rows[0]["method"] == "GET"

    def test_skip_empty_rows(self):
        content = "case_id,method\nC1,GET\n,\nC2,POST\n"
        rows = parse_csv(content)
        assert len(rows) == 2

    def test_auto_parse_json_field(self):
        content = 'case_id,request_body\nC1,"{""title"":""x""}"\n'
        rows = parse_csv(content)
        assert rows[0]["request_body"] == {"title": "x"}

    def test_invalid_json_kept_as_string(self):
        content = 'case_id,request_body\nC1,{not json}\n'
        rows = parse_csv(content)
        assert rows[0]["request_body"] == "{not json}"


class TestParserXLSX:
    def _make_xlsx(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_parse_basic_xlsx(self):
        content = self._make_xlsx([
            ["case_id", "method", "path"],
            ["C1", "GET", "/posts"],
        ])
        rows = parse_xlsx(content)
        assert len(rows) == 1
        assert rows[0]["case_id"] == "C1"
        assert rows[0]["path"] == "/posts"

    def test_xlsx_skips_all_none_row(self):
        content = self._make_xlsx([
            ["case_id", "method"],
            ["C1", "GET"],
            [None, None],
            ["C2", "POST"],
        ])
        rows = parse_xlsx(content)
        assert len(rows) == 2


class TestParseFileDispatch:
    def test_detect_format(self):
        assert detect_format("a.xlsx") == "xlsx"
        assert detect_format("a.csv") == "csv"
        assert detect_format("noext") == "csv"

    def test_parse_file_csv(self):
        content = b"case_id,method\nC1,GET\n"
        rows = parse_file(content, "data.csv")
        assert rows[0]["case_id"] == "C1"

    def test_auto_parse_json_list(self):
        row = {"k": "[1, 2, 3]"}
        assert _auto_parse_json(row)["k"] == [1, 2, 3]


# ==================== renderer ====================

class TestRenderDynamicPlaceholders:
    def test_uuid_replaced(self):
        out = _render_dynamic_placeholders("id={{uuid}}")
        assert "{{uuid}}" not in out and len(out) > 10

    def test_uuid_short_length(self):
        out = _render_dynamic_placeholders("{{uuid_short}}")
        assert len(out) == 8

    def test_timestamp_is_digits(self):
        out = _render_dynamic_placeholders("{{timestamp}}")
        assert out.isdigit()

    def test_random_int_in_range(self):
        out = _render_dynamic_placeholders("{{random_int:5:5}}")
        assert out == "5"

    def test_random_str_length(self):
        out = _render_dynamic_placeholders("{{random_str:12}}")
        assert len(out) == 12

    def test_faker_email_format(self):
        out = _render_dynamic_placeholders("{{faker_email}}")
        assert "@autotest.com" in out

    def test_non_string_returns_unchanged(self):
        assert _render_dynamic_placeholders(123) == 123


class TestRenderAllPlaceholders:
    def test_variable_pool_substitution(self):
        out = _render_all_placeholders("Bearer {{token}}", {"token": "abc"})
        assert out == "Bearer abc"

    def test_nested_dict_and_list(self):
        data = {"h": {"Auth": "{{token}}"}, "arr": ["{{token}}"]}
        out = _render_all_placeholders(data, {"token": "T"})
        assert out["h"]["Auth"] == "T"
        assert out["arr"][0] == "T"

    def test_non_string_passthrough(self):
        assert _render_all_placeholders(42, {}) == 42


# ==================== runner ====================

class TestRunnerBuildRequest:
    def _runner(self):
        return TestRunner(base_url="https://api.example.com", cookie="", timeout=10)

    def test_absolute_url_kept(self):
        r = self._runner()._build_request({"method": "get", "path": "http://x.com/a"})
        assert r["url"] == "http://x.com/a"
        assert r["method"] == "GET"

    def test_relative_path_joined(self):
        r = self._runner()._build_request({"path": "/posts"})
        assert r["url"] == "https://api.example.com/posts"

    def test_path_without_leading_slash(self):
        r = self._runner()._build_request({"path": "posts"})
        assert r["url"] == "https://api.example.com/posts"

    def test_json_body_parsed_from_string(self):
        r = self._runner()._build_request(
            {"path": "/x", "request_body": '{"a": 1}'}
        )
        assert r["json"] == {"a": 1}

    def test_headers_merged(self):
        r = self._runner()._build_request(
            {"path": "/x", "headers": '{"X-Token": "t"}'}
        )
        assert r["headers"]["X-Token"] == "t"
        assert r["headers"]["Content-Type"] == "application/json"


class TestRunnerExtractVars:
    def test_extract_simple(self):
        r = TestRunner("https://x.com", "", 10)
        r._extract_vars("token=data.token", {"data": {"token": "abc"}})
        assert r.pool["token"] == "abc"

    def test_extract_list_index(self):
        r = TestRunner("https://x.com", "", 10)
        r._extract_vars("id=data.list[0].id", {"data": {"list": [{"id": 9}]}})
        assert r.pool["id"] == 9

    def test_extract_missing_ignored(self):
        r = TestRunner("https://x.com", "", 10)
        r._extract_vars("x=data.nope", {"data": {}})
        assert "x" not in r.pool


class _FakeResp:
    def __init__(self, status_code):
        self.status_code = status_code


class TestRunnerAssert:
    def _runner(self):
        return TestRunner("https://x.com", "", 10)

    def test_status_pass(self):
        passed, reason = self._runner()._assert(
            {"expected_status": "200"}, _FakeResp(200), {}, 50
        )
        assert passed and reason == ""

    def test_status_fail(self):
        passed, reason = self._runner()._assert(
            {"expected_status": "200"}, _FakeResp(500), {}, 50
        )
        assert not passed and "状态码" in reason

    def test_max_ms_fail(self):
        passed, reason = self._runner()._assert(
            {"expected_max_ms": "100"}, _FakeResp(200), {}, 500
        )
        assert not passed and "超时" in reason

    def test_field_match(self):
        passed, _ = self._runner()._assert(
            {"expected_title": "hello"}, _FakeResp(200), {"title": "hello"}, 10
        )
        assert passed

    def test_field_mismatch(self):
        passed, reason = self._runner()._assert(
            {"expected_title": "hello"}, _FakeResp(200), {"title": "world"}, 10
        )
        assert not passed and "title" in reason


class TestTruncate:
    def test_short_kept(self):
        assert _truncate({"a": 1}) == '{"a": 1}'

    def test_long_truncated(self):
        out = _truncate("x" * 5000, max_len=100)
        assert out.endswith("...(已截断)")
        assert len(out) < 200
